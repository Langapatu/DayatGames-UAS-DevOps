from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "docs" / "report"
DOCX = OUT / "Laporan_Akhir_DayatGames_Galang_Rispai.docx"

TITLE = (
    "IMPLEMENTASI DEVOPS DAN PENGEMBANGAN AGILE PADA APLIKASI "
    "PENJUALAN GAME DIGITAL DAYATGAMES BERBASIS LARAVEL DAN DOCKER"
)

BLUE = "1E3A5F"
LIGHT_BLUE = "DCE6F1"
LIGHT_GRAY = "F2F2F2"
INK = RGBColor(0, 0, 0)

WORD_TOC_ENTRIES = (
    ("LEMBAR IDENTITAS", 2, 1),
    ("ABSTRAK", 3, 1),
    ("DAFTAR ISI", 4, 1),
    ("DAFTAR GAMBAR", 6, 1),
    ("DAFTAR TABEL", 8, 1),
    ("BAB I - PENDAHULUAN", 9, 1),
    ("1.1 Latar Belakang", 9, 2),
    ("1.2 Rumusan Masalah", 9, 2),
    ("1.3 Tujuan", 9, 2),
    ("1.4 Batasan Masalah", 10, 2),
    ("BAB II - LANDASAN TEORI", 11, 1),
    ("2.1 Agile, Scrum, Backlog, dan SRS", 11, 2),
    ("2.2 Containerization dan Docker Compose", 11, 2),
    ("2.3 Basis Data Relasional", 11, 2),
    ("2.4 Laravel, MySQL, Nginx/PHP-FPM, dan phpMyAdmin", 11, 2),
    ("BAB III - ANALISIS DAN PERANCANGAN", 13, 1),
    ("3.1 Analisis Kebutuhan", 13, 2),
    ("3.2 Perancangan Basis Data", 13, 2),
    ("3.3 Arsitektur Sistem", 15, 2),
    ("BAB IV - IMPLEMENTASI", 16, 1),
    ("4.1 Implementasi Lingkungan Docker", 16, 2),
    ("4.1.1 Struktur Folder", 16, 3),
    ("4.1.2 Dockerfile", 16, 3),
    ("4.1.3 compose.yaml", 16, 3),
    ("4.1.4 Koneksi Aplikasi ke Database", 16, 3),
    ("4.2 Implementasi Basis Data", 18, 2),
    ("4.3 Implementasi CRUD", 21, 2),
    ("4.4 Marketplace dan Transaksi", 25, 2),
    ("4.5 Antarmuka, Responsivitas, dan Motion", 28, 2),
    ("4.6 Version Control", 31, 2),
    ("BAB V - PENGUJIAN", 32, 1),
    ("5.1 Pengujian Fungsional CRUD dan Transaksi", 32, 2),
    ("5.2 Pengujian Koneksi Antar-Service", 32, 2),
    ("5.3 Pengujian Ketahanan dan Persistensi", 34, 2),
    ("BAB VI - KENDALA DAN PENYELESAIAN", 36, 1),
    ("BAB VII - PENUTUP", 37, 1),
    ("7.1 Kesimpulan", 37, 2),
    ("7.2 Saran Pengembangan Selanjutnya", 37, 2),
    ("DAFTAR PUSTAKA", 38, 1),
    ("LAMPIRAN", 39, 1),
    ("Lampiran A. Konfigurasi Infrastruktur", 39, 2),
    ("A.1 Dockerfile lengkap", 39, 2),
    ("A.2 compose.yaml lengkap", 40, 2),
    ("A.3 Konfigurasi Nginx lengkap", 42, 2),
    ("Lampiran B. Diagram Ukuran Penuh", 43, 2),
    ("Lampiran C. Bukti Aplikasi dan Data", 44, 2),
    ("Lampiran C.1 Bukti Redesign dan Responsivitas", 49, 3),
    ("Lampiran D. Akun Demo dan Petunjuk Menjalankan", 53, 2),
    ("Lampiran E. Daftar Bukti dan Pekerjaan Manual", 53, 2),
)

WORD_FIGURE_ENTRIES = (
    ("Gambar 1. Entity Relationship Diagram DayatGames", 14),
    ("Gambar 2. Arsitektur container dan alur request DayatGames", 15),
    ("Gambar 3. Halaman utama DayatGames pada localhost:8080", 17),
    ("Gambar 4. Tabel DayatGames pada phpMyAdmin", 19),
    ("Gambar 5. Validasi foreign key pada phpMyAdmin", 20),
    ("Gambar 6. Daftar game pada shell admin hasil redesign", 22),
    ("Gambar 7. Form tambah game yang telah dikelompokkan", 22),
    ("Gambar 8. Form edit game pada antarmuka admin baru", 23),
    ("Gambar 9. Implementasi CRUD genre", 23),
    ("Gambar 10. Implementasi CRUD publisher", 24),
    ("Gambar 11. Checkout dengan metode pembayaran simulasi", 26),
    ("Gambar 12. Pembayaran terverifikasi oleh admin", 26),
    ("Gambar 13. Game masuk library customer", 27),
    ("Gambar 14. Moderasi review oleh admin", 27),
    ("Gambar 15. Konsep visual redesign DayatGames", 29),
    ("Gambar 16. Katalog customer responsif", 29),
    ("Gambar 17. Detail game dengan metadata dan aksi pembelian", 30),
    ("Gambar 18. Data order hasil browser pada MySQL nyata", 33),
    ("Gambar 19. Data QA Persistence tetap tersedia setelah restart", 34),
    ("Gambar 20. ERD DayatGames untuk lampiran", 43),
    ("Gambar 21. Arsitektur DayatGames untuk lampiran", 43),
    ("Gambar 22. Dashboard admin", 44),
    ("Gambar 23. Daftar game admin", 44),
    ("Gambar 24. Hasil search dan filter", 45),
    ("Gambar 25. Wishlist customer", 45),
    ("Gambar 26. Cart customer", 46),
    ("Gambar 27. Detail order customer", 46),
    ("Gambar 28. Detail pembayaran", 47),
    ("Gambar 29. Form review pemilik game", 47),
    ("Gambar 30. Data game pada phpMyAdmin", 48),
    ("Gambar 31. Data order hasil alur end-to-end", 48),
    ("Gambar 32. Halaman registrasi customer hasil redesign", 49),
    ("Gambar 33. Pengaturan profil customer modern", 49),
    ("Gambar 34. Katalog 4:5 pada viewport mobile", 50),
    ("Gambar 35. Dashboard admin pada viewport mobile", 51),
    ("Gambar 36. Login pada viewport mobile", 52),
)

WORD_TABLE_ENTRIES = (
    ("Tabel 1. Identitas akademik penyusun", 2),
    ("Tabel 2. Ringkasan kebutuhan sistem", 13),
    ("Tabel 3. Deskripsi 15 tabel bisnis inti", 14),
    ("Tabel 4. Pemetaan port host dan container", 15),
    ("Tabel 5. Riwayat commit bertahap sebelum laporan", 31),
    ("Tabel 6. Ringkasan pengujian fungsional", 32),
    ("Tabel 7. Hasil verifikasi teknis akhir", 34),
    ("Tabel 8. Kendala nyata dan penyelesaian", 36),
    ("Tabel 9. Akun demo lokal", 53),
)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_geometry(table, widths_cm: list[float]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    total_dxa = int(sum(widths_cm) / 2.54 * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for cm_width in widths_cm:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(int(cm_width / 2.54 * 1440)))
        grid.append(grid_col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = Cm(widths_cm[idx])
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(widths_cm[idx] / 2.54 * 1440)))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_font(run, size=12, bold=None, italic=None, color=INK, name="Times New Roman"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    return run


def add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=10)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(paragraph, "PAGE", "1")


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(3)
    section.bottom_margin = Cm(3)
    section.left_margin = Cm(4)
    section.right_margin = Cm(3)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 14, 18, 10),
        ("Heading 2", 12, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = INK
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.first_line_indent = Cm(0)

    for style_name in ("Title", "Subtitle", "Caption"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(12)
        style.paragraph_format.left_indent = Cm(1.25)
        style.paragraph_format.first_line_indent = Cm(-0.5)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        style.paragraph_format.space_after = Pt(3)

    for name in ("Figure Caption", "Table Caption"):
        if name not in [s.name for s in doc.styles]:
            style = doc.styles.add_style(name, 1)
        else:
            style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(10)
        style.font.italic = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    add_page_number(footer.paragraphs[0])

    settings = doc.settings._element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "false")
    settings.append(update)


def para(doc, text="", *, bold=False, italic=False, align=None, indent=True, size=12):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(1.25) if indent else Cm(0)
    if align is not None:
        p.alignment = align
    set_font(p.add_run(text), size=size, bold=bold, italic=italic)
    return p


def index_line(doc, label: str, page: int, level: int = 1):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.7 * (level - 1))
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(15.5),
        WD_TAB_ALIGNMENT.RIGHT,
        WD_TAB_LEADER.DOTS,
    )
    set_font(p.add_run(label), size=10)
    set_font(p.add_run(f"\t{page}"), size=10)
    return p


def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(-0.5)
    set_font(p.add_run(text))
    return p


def heading(doc, text, level=1, new_page=False):
    if new_page:
        doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = Cm(0)
    set_font(p.add_run(text), size=14 if level == 1 else 12, bold=True)
    return p


def caption(doc, text, kind="Figure"):
    style = "Figure Caption" if kind == "Figure" else "Table Caption"
    p = doc.add_paragraph(style=style)
    set_font(p.add_run(text), size=10, italic=True)
    return p


def add_picture(doc, path: Path, caption_text: str, width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    caption(doc, caption_text, "Figure")


def add_table(doc, headers, rows, widths_cm, caption_text):
    caption(doc, caption_text, "Table")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[idx], LIGHT_BLUE)
        p = table.rows[0].cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(str(header)), size=9.5, bold=True)
    set_repeat_table_header(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_font(p.add_run(str(value)), size=9)
    set_table_geometry(table, widths_cm)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_code(doc, title, code):
    heading(doc, title, 2)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0.3)
    p.paragraph_format.right_indent = Cm(0.3)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F5F5F5")
    p_pr.append(shd)
    for i, line in enumerate(code.splitlines()):
        if i:
            p.add_run().add_break()
        set_font(p.add_run(line), name="Consolas", size=8.3)


def cover(doc):
    for _ in range(3):
        para(doc, "", indent=False)
    p = para(doc, "LAPORAN AKHIR", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=14)
    p.paragraph_format.space_after = Pt(18)
    p = para(doc, TITLE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, size=14)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(35)
    para(doc, "Disusun untuk memenuhi tugas UAS mata kuliah", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "DevOps dan Pengembangan Agile", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for _ in range(2):
        para(doc, "", indent=False)
    para(doc, "Disusun oleh:", align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "Galang Rispa'i", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    para(doc, "NPM 237006516020", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    for _ in range(3):
        para(doc, "", indent=False)
    for line in (
        "PROGRAM STUDI SISTEM INFORMASI",
        "FAKULTAS TEKNOLOGI KOMUNIKASI DAN INFORMATIKA",
        "UNIVERSITAS NASIONAL",
        "JAKARTA",
        "2026",
    ):
        para(doc, line, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def front_matter(doc):
    heading(doc, "LEMBAR IDENTITAS", 1, new_page=True)
    rows = [
        ("Nama", "Galang Rispa'i"),
        ("NPM", "237006516020"),
        ("Program Studi", "Sistem Informasi"),
        ("Fakultas", "Fakultas Teknologi Komunikasi dan Informatika"),
        ("Universitas", "Universitas Nasional"),
        ("Kelas", "R.02"),
        ("Mata Kuliah", "DevOps dan Pengembangan Agile"),
        ("Dosen", "Lili Dwi Yulianto, S.Kom., M.Kom."),
    ]
    add_table(doc, ("Atribut", "Keterangan"), rows, [4.2, 11.3], "Tabel 1. Identitas akademik penyusun")
    para(
        doc,
        "Catatan sampul: logo Universitas Nasional tidak tersedia di workspace. "
        "Sampul dibuat berbasis teks agar tidak menggunakan aset institusi dari sumber yang tidak terverifikasi.",
        italic=True,
    )

    heading(doc, "ABSTRAK", 1, new_page=True)
    para(
        doc,
        "DayatGames merupakan aplikasi marketplace penjualan game digital untuk simulasi akademik. "
        "Proyek ini menggabungkan praktik pengembangan Agile berbasis product backlog dengan lingkungan "
        "terkontainerisasi Docker Compose. Aplikasi dibangun menggunakan Laravel 13, PHP-FPM 8.4, Nginx, "
        "MySQL 8.0, phpMyAdmin, Blade, Vite, GSAP, Lenis, dan Swiper. Empat service utama terdiri atas app, "
        "webserver, db, dan phpmyadmin. Basis data memuat 15 tabel bisnis dan 20 foreign key. Fitur yang "
        "diimplementasikan mencakup autentikasi dua role, CRUD katalog, pencarian dan filter, wishlist, "
        "cart, checkout, pembayaran simulasi, verifikasi admin, library, review, moderasi, serta dashboard.",
    )
    para(
        doc,
        "Verifikasi akhir menggunakan MySQL nyata dan browser menghasilkan satu order, satu pembayaran "
        "terverifikasi, satu kepemilikan game, dan satu review terpublikasi. Suite Laravel meluluskan 26 "
        "test dengan 148 assertion. Build Vite mengolah 44 modul dan audit dependency melaporkan nol "
        "kerentanan pada tingkat high. Data uji persistensi tetap tersedia sesudah restart seluruh stack. "
        "Hasil ini menunjukkan bahwa implementasi memenuhi kebutuhan fungsional utama, pemisahan role, "
        "integritas transaksi, dan ketahanan data yang ditetapkan dalam backlog.",
    )
    para(
        doc,
        "Kata kunci: DevOps, Agile, Laravel, Docker Compose, MySQL, marketplace game digital.",
        italic=True,
        indent=False,
    )

    for title, entries in (
        ("DAFTAR ISI", WORD_TOC_ENTRIES),
        ("DAFTAR GAMBAR", tuple((label, page, 1) for label, page in WORD_FIGURE_ENTRIES)),
        ("DAFTAR TABEL", tuple((label, page, 1) for label, page in WORD_TABLE_ENTRIES)),
    ):
        heading(doc, title, 1, new_page=True)
        for label, page, level in entries:
            index_line(doc, label, page, level)


def chapter_one(doc):
    heading(doc, "BAB I - PENDAHULUAN", 1, new_page=True)
    heading(doc, "1.1 Latar Belakang", 2)
    para(
        doc,
        "Aplikasi perdagangan digital tidak cukup dinilai dari tampilan antarmuka. Sistem harus menjaga "
        "konsistensi data, memisahkan hak akses, mengamankan proses transaksi, dan dapat dijalankan secara "
        "berulang pada lingkungan yang dapat direproduksi. DayatGames dirancang untuk menjawab kebutuhan "
        "tersebut melalui satu proyek utuh yang mencakup aplikasi, infrastruktur container, basis data, "
        "pengujian, bukti eksekusi, dan version control.",
    )
    para(
        doc,
        "Praktik Agile diterapkan dengan menerjemahkan kebutuhan menjadi functional requirement, "
        "non-functional requirement, product backlog, acceptance criteria, serta traceability matrix. "
        "Praktik DevOps diterapkan melalui Docker Compose, healthcheck, named volume, pengujian otomatis, "
        "build aset, dokumentasi perintah, dan commit bertahap. Dengan demikian, hasil akhir bukan landing "
        "page, melainkan marketplace yang menjalankan alur data dari katalog sampai kepemilikan game.",
    )
    heading(doc, "1.2 Rumusan Masalah", 2)
    for text in (
        "Bagaimana membangun marketplace game digital berbasis Laravel yang memiliki CRUD, transaksi, dan otorisasi nyata?",
        "Bagaimana menghubungkan Nginx, PHP-FPM/Laravel, MySQL, dan phpMyAdmin dalam satu Docker Compose?",
        "Bagaimana menjaga integritas harga, order, pembayaran, library, dan review pada alur transaksi?",
        "Bagaimana membuktikan fungsi sistem melalui test otomatis, browser, phpMyAdmin, dan uji persistensi?",
    ):
        bullet(doc, text)
    heading(doc, "1.3 Tujuan", 2)
    para(
        doc,
        "Tujuan proyek adalah menghasilkan aplikasi DayatGames yang fungsional, dapat dijalankan melalui "
        "Docker Compose, memiliki skema relasional terdokumentasi, memenuhi backlog prioritas, dan "
        "menyediakan bukti pengujian yang dapat ditelusuri ke requirement.",
    )
    heading(doc, "1.4 Batasan Masalah", 2)
    for text in (
        "Produk merupakan marketplace fungsional, bukan landing page.",
        "CRUD nyata disediakan terutama untuk games, genres, dan publishers; developer menjadi CRUD tambahan.",
        "Basis data menggunakan 15 tabel bisnis, melampaui batas minimum delapan tabel.",
        "Lingkungan runtime menggunakan Docker Compose pada komputer lokal.",
        "Pembayaran adalah simulasi akademik dan bukan integrasi payment gateway atau bank sungguhan.",
        "Harga game merupakan data demonstrasi; 16 harga IDR diperiksa dari Steam wilayah Indonesia pada 24-25 Juli 2026 dan dua harga ditandai demo.",
        "Repository GitHub belum dibuat karena GitHub CLI tidak tersedia; repository lokal dan riwayat commit telah lengkap.",
    ):
        bullet(doc, text)


def chapter_two(doc):
    heading(doc, "BAB II - LANDASAN TEORI", 1, new_page=True)
    heading(doc, "2.1 Agile, Scrum, Backlog, dan SRS", 2)
    para(
        doc,
        "Scrum adalah kerangka kerja ringan untuk menghasilkan nilai melalui solusi adaptif terhadap "
        "masalah kompleks. Panduan Scrum 2020 menjelaskan Product Backlog sebagai daftar terurut mengenai "
        "hal yang diperlukan untuk meningkatkan produk. Dalam proyek ini, backlog berfungsi sebagai "
        "jembatan antara spesifikasi kebutuhan dan increment perangkat lunak; setiap item memiliki "
        "acceptance criteria, status, bukti implementasi, serta test terkait (Schwaber dan Sutherland, 2020).",
    )
    heading(doc, "2.2 Containerization dan Docker Compose", 2)
    para(
        doc,
        "Container mengemas aplikasi bersama dependency runtime sehingga lingkungan lebih konsisten. "
        "Docker Compose mendefinisikan service, network, volume, port, dependency, dan healthcheck dalam "
        "satu berkas deklaratif. Dokumentasi Docker menjelaskan bahwa service pada network Compose dapat "
        "saling ditemukan melalui nama service. Oleh karena itu, aplikasi memakai DB_HOST=db, bukan "
        "127.0.0.1, dan berkomunikasi ke port container 3306 (Docker, 2026a).",
    )
    heading(doc, "2.3 Basis Data Relasional", 2)
    para(
        doc,
        "Basis data relasional mengorganisasi data ke dalam tabel dengan primary key, foreign key, dan "
        "constraint. InnoDB mendukung transaksi dan foreign key untuk menjaga konsistensi data terkait. "
        "Aksi referensial seperti RESTRICT, CASCADE, dan SET NULL dipilih berdasarkan karakter data; data "
        "transaksi dipertahankan, sedangkan data sementara seperti cart dapat dihapus berantai (Oracle, 2026).",
    )
    heading(doc, "2.4 Laravel, MySQL, Nginx/PHP-FPM, dan phpMyAdmin", 2)
    para(
        doc,
        "Laravel menyediakan routing, middleware, validation, ORM Eloquent, migration, seeder, database "
        "transaction, dan fasilitas testing. Eloquent memetakan model ke tabel serta mendukung operasi "
        "insert, update, delete, dan relationship. Nginx melayani aset statis dan meneruskan request PHP "
        "melalui FastCGI ke PHP-FPM. MySQL menyimpan data utama, sedangkan phpMyAdmin menyediakan antarmuka "
        "administrasi database pada jaringan yang sama (Laravel, 2026; Nginx, 2026; phpMyAdmin, 2026).",
    )


def chapter_three(doc):
    heading(doc, "BAB III - ANALISIS DAN PERANCANGAN", 1, new_page=True)
    heading(doc, "3.1 Analisis Kebutuhan", 2)
    para(
        doc,
        "Analisis menghasilkan 20 functional requirement dan 12 non-functional requirement. Aktor terdiri "
        "atas guest, customer, dan admin. Guest dapat membaca katalog; customer mengelola wishlist, cart, "
        "checkout, order, library, profil, dan review; admin mengelola master katalog, transaksi, customer, "
        "pembayaran, review, dan statistik. Seluruh requirement dipetakan ke backlog, route/controller/view, "
        "tabel, test, serta screenshot dalam traceability matrix.",
    )
    requirements = [
        ("Guest", "Home, katalog published, search/filter, detail, registrasi, login"),
        ("Customer", "Profil, wishlist, cart, checkout, order, payment, library, review"),
        ("Admin", "Dashboard, CRUD katalog, customer, order, payment, review, featured"),
        ("NFR", "Compose, MySQL, Nginx/PHP-FPM, phpMyAdmin, keamanan, responsive, persistensi, Git"),
    ]
    add_table(doc, ("Aktor/jenis", "Ruang lingkup"), requirements, [3.2, 12.3], "Tabel 2. Ringkasan kebutuhan sistem")

    heading(doc, "3.2 Perancangan Basis Data", 2)
    para(
        doc,
        "Skema inti memiliki 15 tabel bisnis dan 20 foreign key. Order item menyimpan snapshot judul, "
        "harga awal, nilai diskon, dan subtotal agar histori tidak berubah ketika master game diedit. "
        "Constraint unik pada cart_items, wishlists, payments, libraries, dan reviews mencegah duplikasi. "
        "Pembatasan penghapusan pada orders dan libraries mempertahankan riwayat transaksi.",
    )
    add_picture(doc, ROOT / "docs" / "ERD.png", "Gambar 1. Entity Relationship Diagram DayatGames", 15.5)
    tables = [
        ("users", "Akun dan role", "email unik; role admin/customer"),
        ("developers", "Master pengembang", "slug unik; delete restricted"),
        ("publishers", "Master penerbit", "slug unik; delete restricted"),
        ("genres", "Master genre", "slug unik"),
        ("games", "Master produk game", "FK developer/publisher; slug unik"),
        ("game_genre", "Pivot game-genre", "PK gabungan"),
        ("game_images", "Galeri game", "FK game; cascade"),
        ("carts", "Keranjang per user", "user_id unik"),
        ("cart_items", "Isi keranjang", "cart_id + game_id unik"),
        ("wishlists", "Daftar keinginan", "user_id + game_id unik"),
        ("orders", "Header transaksi", "order_code unik; delete restricted"),
        ("order_items", "Snapshot item transaksi", "game_id nullable; histori tetap"),
        ("payments", "Pembayaran simulasi", "order_id unik; status"),
        ("libraries", "Kepemilikan game", "user_id + game_id unik"),
        ("reviews", "Ulasan pembeli", "rating 1-5; satu per user/game"),
    ]
    add_table(doc, ("Tabel", "Fungsi", "Constraint utama"), tables, [3.2, 5.0, 7.3], "Tabel 3. Deskripsi 15 tabel bisnis inti")

    heading(doc, "3.3 Arsitektur Sistem", 2)
    para(
        doc,
        "Browser mengakses host localhost:8080 yang dipetakan ke port 80 container webserver. Nginx "
        "meneruskan request dinamis ke app:9000 melalui FastCGI. Laravel mengakses MySQL melalui db:3306. "
        "Akses administrasi database memakai localhost:8081 menuju port 80 container phpmyadmin yang juga "
        "terhubung ke db:3306. Source code dipasang sebagai bind mount, sedangkan data MySQL disimpan pada "
        "named volume dayatgames_db_data.",
    )
    add_picture(doc, ROOT / "docs" / "ARCHITECTURE.png", "Gambar 2. Arsitektur container dan alur request DayatGames", 15.5)
    ports = [
        ("8080:80", "Host 8080", "Nginx 80", "Aplikasi web"),
        ("3306:3306", "Host 3306", "MySQL 3306", "Akses database dari host"),
        ("8081:80", "Host 8081", "phpMyAdmin 80", "Administrasi database"),
    ]
    add_table(doc, ("Mapping", "Sisi kiri", "Sisi kanan", "Fungsi"), ports, [3.0, 3.5, 4.0, 5.0], "Tabel 4. Pemetaan port host dan container")


def chapter_four(doc):
    heading(doc, "BAB IV - IMPLEMENTASI", 1, new_page=True)
    heading(doc, "4.1 Implementasi Lingkungan Docker", 2)
    heading(doc, "4.1.1 Struktur Folder", 3)
    para(
        doc,
        "Project final ditempatkan pada C:\\.Kuliah\\TugasMatkul\\DevOPS\\UAS\\DayatGames. Project lama "
        "ditemukan melalui label dan mount container pada C:\\Coding\\laravel-docker, kemudian hanya "
        "diinspeksi. Struktur final memisahkan app, database, docker/nginx, public, resources, routes, "
        "tests, dan docs. Lima screenshot editor/terminal yang tidak dapat diambil oleh browser dicatat "
        "sebagai pekerjaan manual di SCREENSHOT_CHECKLIST.md dan tidak disintesis.",
    )
    heading(doc, "4.1.2 Dockerfile", 3)
    para(
        doc,
        "Dockerfile memakai php:8.4-fpm. Instruksi RUN memasang extension pdo_mysql, mbstring, gd, zip, "
        "dan dependency build. COPY dari image composer menyediakan binary Composer. WORKDIR menetapkan "
        "/var/www, COPY memindahkan manifest dependency dan source, sedangkan CMD menjalankan php-fpm.",
    )
    heading(doc, "4.1.3 compose.yaml", 3)
    para(
        doc,
        "Service app dibangun dari Dockerfile dan menunggu db sehat. Service webserver memakai Nginx "
        "Alpine dan memetakan 8080:80. Service db memakai MySQL 8.0, healthcheck, serta named volume. "
        "Service phpmyadmin memetakan 8081:80 dan memakai PMA_HOST=db. Restart policy unless-stopped "
        "diterapkan pada seluruh service.",
    )
    heading(doc, "4.1.4 Koneksi Aplikasi ke Database", 3)
    para(
        doc,
        "DB_HOST harus bernilai db karena 127.0.0.1 di dalam container app menunjuk kembali ke container "
        "app, bukan MySQL. Docker DNS mendaftarkan nama service db pada network dayatgames_network. Port "
        "kiri pada mapping adalah port host, sedangkan port kanan adalah port container. Kredensial "
        "DB_DATABASE, DB_USERNAME, dan DB_PASSWORD pada .env Laravel disamakan dengan MYSQL_DATABASE, "
        "MYSQL_USER, dan MYSQL_PASSWORD di Compose.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "05-aplikasi-home.png", "Gambar 3. Halaman utama DayatGames pada localhost:8080", 12.6)

    heading(doc, "4.2 Implementasi Basis Data", 2, new_page=True)
    para(
        doc,
        "Migration katalog membentuk master game dan relasi genre; migration commerce membentuk cart, "
        "wishlist, order, payment, library, dan review. Model Eloquent mendefinisikan belongsTo, hasMany, "
        "hasOne, dan belongsToMany. Seeder menggunakan updateOrCreate sehingga dapat dijalankan ulang "
        "tanpa menggandakan data master. Validasi schema dilakukan dengan test, INFORMATION_SCHEMA, dan "
        "phpMyAdmin yang terhubung ke database dayatgames yang sama.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "28-phpmyadmin-tables.png", "Gambar 4. Tabel DayatGames pada phpMyAdmin", 14.2)
    add_picture(doc, ROOT / "docs" / "screenshots" / "29-foreign-keys.png", "Gambar 5. Validasi foreign key pada phpMyAdmin", 14.7)

    heading(doc, "4.3 Implementasi CRUD", 2, new_page=True)
    para(
        doc,
        "CRUD admin memakai resource route, route model binding, Form Request, middleware role, CSRF, "
        "validasi upload, pagination, pencarian, flash message, dan konfirmasi delete. Operasi browser "
        "untuk game diuji dengan membuat QA Browser Game, mengubah nama dan harga, kemudian menghapusnya. "
        "Genre dan publisher juga memiliki form create/update serta proteksi penghapusan saat direferensikan.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "09-games-index.png", "Gambar 6. Daftar game pada shell admin hasil redesign", 13.3)
    add_picture(doc, ROOT / "docs" / "screenshots" / "10-create-game.png", "Gambar 7. Form tambah game yang telah dikelompokkan", 13.3)
    add_picture(doc, ROOT / "docs" / "screenshots" / "12-edit-game.png", "Gambar 8. Form edit game pada antarmuka admin baru", 12.6)
    add_picture(doc, ROOT / "docs" / "screenshots" / "15-crud-genre.png", "Gambar 9. Implementasi CRUD genre", 14.8)
    add_picture(doc, ROOT / "docs" / "screenshots" / "16-crud-publisher.png", "Gambar 10. Implementasi CRUD publisher", 14.0)

    heading(doc, "4.4 Marketplace dan Transaksi", 2, new_page=True)
    para(
        doc,
        "Harga checkout selalu dihitung ulang dari database, sehingga input harga dari browser tidak "
        "dipercaya. Pembuatan order, order_items, dan payment dilakukan dalam database transaction. "
        "Verifikasi admin juga transaksional dan idempotent: payment menjadi verified, order completed, "
        "game masuk library melalui operasi yang tidak menggandakan pasangan user-game, dan cart terkait "
        "dibersihkan. Hanya pemilik library yang dapat mengirim review.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "22-checkout.png", "Gambar 11. Checkout dengan metode pembayaran simulasi", 14.0)
    add_picture(doc, ROOT / "docs" / "screenshots" / "25-payment-verified.png", "Gambar 12. Pembayaran terverifikasi oleh admin", 14.7)
    add_picture(doc, ROOT / "docs" / "screenshots" / "26-library.png", "Gambar 13. Game masuk library customer", 14.2)
    add_picture(doc, ROOT / "docs" / "screenshots" / "27-review-moderation.png", "Gambar 14. Moderasi review oleh admin", 14.7)

    heading(doc, "4.5 Antarmuka, Responsivitas, dan Motion", 2, new_page=True)
    para(
        doc,
        "Antarmuka memakai palet gelap premium, Blade components, navbar customer yang ringkas, autentikasi "
        "split-panel, sidebar admin, tabel dan form terkelompok, badge, empty state, serta focus state. "
        "Artwork game memakai media frame konsisten dan object-fit contain agar rasio potret, persegi, maupun "
        "landscape tetap utuh tanpa terpotong. GSAP dan ScrollTrigger menangani reveal/parallax, Lenis menangani "
        "smooth scroll, dan Swiper menangani carousel. prefers-reduced-motion menonaktifkan motion besar "
        "dan smooth scroll. Fitur inti tetap menggunakan form dan link server-side sehingga berfungsi jika "
        "JavaScript animasi gagal.",
    )
    add_picture(doc, ROOT / "docs" / "design" / "dayatgames-ui-concept-v2.png", "Gambar 15. Konsep visual redesign DayatGames", 14.5)
    add_picture(doc, ROOT / "docs" / "screenshots" / "17-katalog-customer.png", "Gambar 16. Katalog customer responsif", 12.4)
    add_picture(doc, ROOT / "docs" / "screenshots" / "18-detail-game.png", "Gambar 17. Detail game dengan metadata dan aksi pembelian", 12.8)

    heading(doc, "4.6 Version Control", 2, new_page=True)
    para(
        doc,
        "Repository lokal menggunakan branch feature/dayatgames-uas. Perubahan dipisahkan menjadi commit "
        "inspeksi, spesifikasi, Docker, database, autentikasi, admin, customer, transaksi, antarmuka, QA, "
        "dan iterasi redesign berbasis hasil evaluasi pengguna. "
        "File .env, vendor, dan node_modules diabaikan. Pemeriksaan pola token/private key pada source "
        "terlacak tidak menemukan rahasia. URL repository belum dicantumkan karena gh tidak tersedia dan "
        "tidak ada push yang dapat dibuktikan.",
    )
    commits = [
        ("639acf9", "chore: inspect legacy docker project and initialize workspace"),
        ("80fd337", "docs: add dayatgames specification and backlog"),
        ("6aea53e", "chore: configure docker compose environment"),
        ("9934dd3", "feat: add database schema and eloquent relationships"),
        ("f878e42", "feat: add authentication and role authorization"),
        ("82f56d7", "feat: implement admin catalog management"),
        ("dd47248", "feat: build customer marketplace and cart"),
        ("402f288", "feat: implement transaction and library workflow"),
        ("6cfeba8", "feat: add responsive interface and motion"),
        ("29a3952", "test: document end-to-end QA evidence"),
        ("4dd0f87", "feat: redesign auth navigation and admin ui"),
    ]
    add_table(doc, ("Commit", "Pesan"), commits, [3.0, 12.5], "Tabel 5. Riwayat commit bertahap sebelum laporan")


def chapter_five(doc):
    heading(doc, "BAB V - PENGUJIAN", 1, new_page=True)
    heading(doc, "5.1 Pengujian Fungsional CRUD dan Transaksi", 2)
    para(
        doc,
        "Pengujian otomatis dijalankan di container app. Hasil akhir adalah 26 test lulus dengan 148 "
        "assertion. Cakupan meliputi autentikasi, role, skema, CRUD master, validasi diskon, katalog, "
        "wishlist, cart, checkout, snapshot harga, otorisasi order, verifikasi dan penolakan pembayaran, "
        "idempotensi library, serta review.",
    )
    tests = [
        ("Katalog", "Buka published/draft", "Hanya published tampil", "Published tampil; draft 404", "Lulus"),
        ("Auth/role", "Akses admin 3 role", "Redirect/403/200", "Sesuai role", "Lulus"),
        ("CRUD", "Create/update/delete", "Data dan relasi berubah", "4 test, 29 assertion", "Lulus"),
        ("Wishlist/cart", "Add dua kali/remove", "Tidak duplikat", "Unique terjaga", "Lulus"),
        ("Checkout", "Kirim price=1", "Pakai harga DB", "Total memakai snapshot DB", "Lulus"),
        ("Order privacy", "User lain buka order", "Ditolak", "HTTP 404", "Lulus"),
        ("Verify", "Verify dua kali", "Satu library", "Idempotent", "Lulus"),
        ("Review", "Owner/non-owner", "403/berhasil", "Sesuai", "Lulus"),
    ]
    add_table(doc, ("Fitur", "Langkah", "Harapan", "Aktual", "Status"), tests, [2.4, 3.6, 3.5, 4.2, 1.8], "Tabel 6. Ringkasan pengujian fungsional")

    heading(doc, "5.2 Pengujian Koneksi Antar-Service", 2)
    para(
        doc,
        "docker compose config --quiet dan build berhasil. Keempat service berstatus Up, sedangkan db "
        "berstatus healthy. Setelah restart, endpoint /, /games, /login, dan phpMyAdmin mengembalikan HTTP "
        "200. Login phpMyAdmin memakai user aplikasi memperlihatkan database, tabel, foreign key, dan data "
        "yang sama dengan aplikasi.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "30b-phpmyadmin-orders.png", "Gambar 18. Data order hasil browser pada MySQL nyata", 14.7)

    heading(doc, "5.3 Pengujian Ketahanan dan Persistensi", 2, new_page=True)
    para(
        doc,
        "Genre QA Persistence dibuat dengan ID 13 sebelum restart. Seluruh service kemudian direstart. "
        "Query pertama terjadi saat MySQL belum sehat dan menghasilkan connection refused; setelah "
        "healthcheck berubah menjadi healthy, query diulang dan menemukan record yang sama. Uji ini "
        "membuktikan data tersimpan pada named volume, bukan lapisan sementara container.",
    )
    add_picture(doc, ROOT / "docs" / "screenshots" / "32-data-after-restart.png", "Gambar 19. Data QA Persistence tetap tersedia setelah restart", 14.5)
    results = [
        ("docker compose config --quiet", "Berhasil"),
        ("docker compose build", "Berhasil setelah public/storage dikecualikan"),
        ("docker compose ps", "4 service Up; db healthy"),
        ("php artisan test", "26 test; 148 assertion; lulus"),
        ("npm run build", "44 modul; berhasil"),
        ("npm audit --audit-level=high", "0 vulnerability"),
        ("php artisan route:list", "60 route"),
        ("Restart dan HTTP", "Data persisten; 4 endpoint HTTP 200"),
    ]
    add_table(doc, ("Perintah/cek", "Hasil aktual"), results, [7.0, 8.5], "Tabel 7. Hasil verifikasi teknis akhir")


def chapter_six_seven(doc):
    heading(doc, "BAB VI - KENDALA DAN PENYELESAIAN", 1, new_page=True)
    issues = [
        (
            "Lokasi project Docker lama",
            "Lokasi awal tidak diasumsikan. Metadata container dan label Compose menunjukkan C:\\Coding\\laravel-docker. Project hanya dibaca dan hasil final dibuat terpisah.",
        ),
        (
            "Dependency Composer pada image",
            "Dependency development yang dipakai command aplikasi perlu tersedia. Dockerfile mempertahankan instalasi dependency yang kompatibel dengan kebutuhan UAS.",
        ),
        (
            "Remote font dan build deterministik",
            "Referensi font remote dihapus agar build tidak bergantung pada jaringan eksternal.",
        ),
        (
            "Cast NUMERIC pada MySQL",
            "Cast NUMERIC tidak kompatibel dengan kebutuhan model; diganti menjadi DECIMAL(15,2) yang sesuai nominal uang.",
        ),
        (
            "Windows public/storage reparse point",
            "Docker build sempat gagal dengan invalid file request. public/storage dikecualikan melalui .dockerignore karena symlink dibuat pada runtime.",
        ),
        (
            "Timing restart MySQL",
            "Query pertama terlalu cepat dan mendapat connection refused. Verifikasi diulang setelah healthcheck db menunjukkan healthy.",
        ),
        (
            "Keterbatasan screenshot terminal/editor",
            "Browser automation tidak dapat mengambil lima target editor/terminal. Target tersebut ditandai Manual secara spesifik dan tidak dibuat sintetis.",
        ),
        (
            "GitHub CLI tidak tersedia",
            "Repository lokal diselesaikan dan diperiksa. URL tidak diklaim; perintah push manual disiapkan pada README.",
        ),
    ]
    add_table(doc, ("Kendala", "Penyelesaian aktual"), issues, [4.6, 10.9], "Tabel 8. Kendala nyata dan penyelesaian")

    heading(doc, "BAB VII - PENUTUP", 1, new_page=True)
    heading(doc, "7.1 Kesimpulan", 2)
    para(
        doc,
        "DayatGames berhasil diimplementasikan sebagai marketplace game digital fungsional berbasis "
        "Laravel dan Docker Compose. Sistem menggunakan empat service terintegrasi, 15 tabel bisnis, "
        "20 foreign key, autentikasi dua role, CRUD lengkap, katalog, transaksi simulasi, library, review, "
        "serta pengujian otomatis dan browser. Named volume menjaga persistensi data, sedangkan healthcheck "
        "mengendalikan dependency startup. Traceability dari requirement sampai test dan bukti tersedia "
        "dalam dokumentasi proyek.",
    )
    para(
        doc,
        "Keberhasilan teknis didukung oleh hasil 26 test dan 148 assertion, build 44 modul, nol vulnerability "
        "tingkat high, 60 route, browser flow MySQL nyata, dan uji restart. Pembayaran, harga, dan repository "
        "dilaporkan sesuai batas faktual: payment adalah simulasi akademik, dua harga merupakan demo, serta "
        "URL GitHub belum tersedia.",
    )
    heading(doc, "7.2 Saran Pengembangan Selanjutnya", 2)
    for item in (
        "Menambahkan object storage dan pipeline optimasi gambar untuk deployment produksi.",
        "Mengintegrasikan payment gateway sandbox dengan webhook terverifikasi.",
        "Menambahkan CI/CD, code coverage, static analysis, dan deployment staging.",
        "Menambahkan audit log administratif dan notifikasi transaksi.",
        "Melakukan pengujian aksesibilitas dan performa pada beberapa browser/perangkat fisik.",
        "Mengunggah repository setelah verifikasi akun GitHub dan melengkapi lima screenshot editor/terminal manual.",
    ):
        bullet(doc, item)


def bibliography(doc):
    heading(doc, "DAFTAR PUSTAKA", 1, new_page=True)
    refs = [
        "Docker, Inc. (2026a). Networking in Compose. https://docs.docker.com/compose/how-tos/networking/ (diakses 24 Juli 2026).",
        "Docker, Inc. (2026b). Compose file reference: Services. https://docs.docker.com/reference/compose-file/services/ (diakses 24 Juli 2026).",
        "Laravel LLC. (2026). Laravel 13.x Documentation: Eloquent, Database, and Testing. https://laravel.com/docs/13.x (diakses 24 Juli 2026).",
        "Nginx. (2026). Beginner's Guide: FastCGI Proxying. https://nginx.org/en/docs/beginners_guide.html (diakses 24 Juli 2026).",
        "Oracle Corporation. (2026). MySQL 8.0 Reference Manual: InnoDB and Foreign Key Constraints. https://dev.mysql.com/doc/refman/8.0/en/create-table-foreign-keys.html (diakses 24 Juli 2026).",
        "phpMyAdmin contributors. (2026). phpMyAdmin Documentation: Installation and Docker Configuration. https://docs.phpmyadmin.net/en/latest/setup.html (diakses 24 Juli 2026).",
        "Schwaber, K., & Sutherland, J. (2020). The Scrum Guide. https://scrumguides.org/docs/scrumguide/v2020/2020-Scrum-Guide-US.pdf.",
    ]
    for ref in refs:
        p = para(doc, ref, indent=False)
        p.paragraph_format.left_indent = Cm(1.25)
        p.paragraph_format.first_line_indent = Cm(-1.25)


def appendices(doc):
    heading(doc, "LAMPIRAN", 1, new_page=True)
    heading(doc, "Lampiran A. Konfigurasi Infrastruktur", 2)
    add_code(doc, "A.1 Dockerfile lengkap", (ROOT / "Dockerfile").read_text(encoding="utf-8"))
    doc.add_page_break()
    add_code(doc, "A.2 compose.yaml lengkap", (ROOT / "compose.yaml").read_text(encoding="utf-8"))
    doc.add_page_break()
    add_code(doc, "A.3 Konfigurasi Nginx lengkap", (ROOT / "docker" / "nginx" / "default.conf").read_text(encoding="utf-8"))

    heading(doc, "Lampiran B. Diagram Ukuran Penuh", 2, new_page=True)
    add_picture(doc, ROOT / "docs" / "ERD.png", "Gambar 20. ERD DayatGames untuk lampiran", 15.5)
    add_picture(doc, ROOT / "docs" / "ARCHITECTURE.png", "Gambar 21. Arsitektur DayatGames untuk lampiran", 15.5)

    heading(doc, "Lampiran C. Bukti Aplikasi dan Data", 2, new_page=True)
    for filename, title, width in (
        ("08-dashboard-admin.png", "Gambar 22. Dashboard admin", 14.0),
        ("09-games-index.png", "Gambar 23. Daftar game admin", 12.8),
        ("19-search-filter.png", "Gambar 24. Hasil search dan filter", 14.1),
        ("20-wishlist.png", "Gambar 25. Wishlist customer", 14.2),
        ("21-cart.png", "Gambar 26. Cart customer", 14.8),
        ("23-order-detail.png", "Gambar 27. Detail order customer", 14.8),
        ("24-payment-detail.png", "Gambar 28. Detail pembayaran", 14.8),
        ("27-review-form.png", "Gambar 29. Form review pemilik game", 14.3),
        ("30a-phpmyadmin-games.png", "Gambar 30. Data game pada phpMyAdmin", 14.0),
        ("30-data-order.png", "Gambar 31. Data order hasil alur end-to-end", 14.8),
    ):
        add_picture(doc, ROOT / "docs" / "screenshots" / filename, title, width)

    heading(doc, "Lampiran C.1 Bukti Redesign dan Responsivitas", 3, new_page=True)
    for filename, title, width in (
        ("33-register-customer.png", "Gambar 32. Halaman registrasi customer hasil redesign", 14.5),
        ("37-profile-customer.png", "Gambar 33. Pengaturan profil customer modern", 14.5),
        ("34-responsive-catalog-mobile.png", "Gambar 34. Katalog 4:5 pada viewport mobile", 6.8),
        ("35-responsive-admin-mobile.png", "Gambar 35. Dashboard admin pada viewport mobile", 6.8),
        ("36-responsive-login-mobile.png", "Gambar 36. Login pada viewport mobile", 6.8),
    ):
        add_picture(doc, ROOT / "docs" / "screenshots" / filename, title, width)

    heading(doc, "Lampiran D. Akun Demo dan Petunjuk Menjalankan", 2, new_page=True)
    accounts = [
        ("Admin", "admin@dayatgames.test", "password", "Akses seluruh admin"),
        ("Customer", "customer@dayatgames.test", "password", "Akses transaksi customer"),
    ]
    add_table(doc, ("Role", "Email", "Password", "Keterangan"), accounts, [2.4, 5.5, 2.8, 4.8], "Tabel 9. Akun demo lokal")
    for step in (
        "Salin .env.example menjadi .env dan pastikan kredensial database konsisten dengan Compose.",
        "Jalankan docker compose config --quiet dan docker compose build.",
        "Jalankan docker compose up -d lalu pastikan db healthy melalui docker compose ps.",
        "Jalankan docker compose exec -T app php artisan key:generate.",
        "Jalankan docker compose exec -T app php artisan migrate --seed.",
        "Jalankan docker compose exec -T app php artisan storage:link.",
        "Jalankan npm install dan npm run build jika aset belum tersedia.",
        "Buka http://localhost:8080 dan http://localhost:8081.",
        "Jalankan docker compose exec -T app php artisan test untuk verifikasi.",
    ):
        bullet(doc, step)
    para(
        doc,
        "Repository: belum memiliki URL remote yang dapat diverifikasi. Repository lokal berada pada "
        "branch feature/dayatgames-uas. Setelah autentikasi GitHub tersedia, buat remote "
        "DayatGames-UAS-DevOps, periksa kembali rahasia, lalu push branch.",
        italic=True,
    )
    heading(doc, "Lampiran E. Daftar Bukti dan Pekerjaan Manual", 2)
    para(
        doc,
        "Folder docs/screenshots berisi 36 screenshot nyata aplikasi dan phpMyAdmin. Lima bukti berikut "
        "belum tersedia karena memerlukan editor/terminal desktop: 01 struktur folder, 02 Dockerfile, "
        "03 compose.yaml, 04 docker compose ps, dan 31 git log. Langkah pengambilan yang spesifik tersedia "
        "pada docs/SCREENSHOT_CHECKLIST.md. Ketidakhadiran bukti ini tidak diganti dengan gambar sintetis.",
    )


def add_core_properties(doc):
    doc.core_properties.title = TITLE
    doc.core_properties.subject = "Laporan UAS DevOps dan Pengembangan Agile"
    doc.core_properties.author = "Galang Rispa'i"
    doc.core_properties.keywords = "DayatGames, Laravel, Docker, DevOps, Agile"
    doc.core_properties.comments = "Dibangun dari implementasi dan bukti QA aktual proyek DayatGames."


def validate_docx(path: Path):
    with ZipFile(path) as z:
        names = set(z.namelist())
        required = {"word/document.xml", "word/styles.xml", "word/settings.xml"}
        missing = required - names
        if missing:
            raise RuntimeError(f"DOCX tidak lengkap: {sorted(missing)}")
        document_xml = z.read("word/document.xml").decode("utf-8")
        for token in ("BAB I", "BAB VII", "DAFTAR PUSTAKA", "LAMPIRAN"):
            if token not in document_xml:
                raise RuntimeError(f"Bagian wajib tidak ditemukan: {token}")

    rendered = Document(path)
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    figure_captions = [
        paragraph.text
        for paragraph in rendered.paragraphs
        if paragraph.style and paragraph.style.name == "Figure Caption"
    ]
    table_captions = [
        paragraph.text
        for paragraph in rendered.paragraphs
        if paragraph.style and paragraph.style.name == "Table Caption"
    ]
    if figure_captions != [label for label, _ in WORD_FIGURE_ENTRIES]:
        raise RuntimeError("Daftar gambar statis tidak sama dengan caption aktual.")
    if table_captions != [label for label, _ in WORD_TABLE_ENTRIES]:
        raise RuntimeError("Daftar tabel statis tidak sama dengan caption aktual.")
    if len(rendered.inline_shapes) != len(WORD_FIGURE_ENTRIES):
        raise RuntimeError("Jumlah gambar inline tidak sesuai daftar gambar.")
    if "Daftar akan diperbarui otomatis saat dokumen dibuka." in text or "\t?" in text:
        raise RuntimeError("Placeholder indeks masih ditemukan.")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_core_properties(doc)
    cover(doc)
    front_matter(doc)
    chapter_one(doc)
    chapter_two(doc)
    chapter_three(doc)
    chapter_four(doc)
    chapter_five(doc)
    chapter_six_seven(doc)
    bibliography(doc)
    appendices(doc)
    doc.save(DOCX)
    validate_docx(DOCX)
    print(f"CREATED={DOCX}")
    print(f"SIZE={DOCX.stat().st_size}")


if __name__ == "__main__":
    main()

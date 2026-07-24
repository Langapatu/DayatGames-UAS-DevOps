from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def add_bookmark(paragraph, name: str, bookmark_id: int):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def field_run(instruction: str, placeholder: str):
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.extend([fonts, size])
    run.append(r_pr)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.extend([begin, instr, separate, text, end])
    return run


def index_paragraph(text: str, bookmark: str):
    p = OxmlElement("w:p")
    p_pr = OxmlElement("w:pPr")
    p_style = OxmlElement("w:pStyle")
    p_style.set(qn("w:val"), "TOC1")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:leader"), "dot")
    tab.set(qn("w:pos"), "8780")
    tabs.append(tab)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "40")
    p_pr.extend([p_style, tabs, spacing])
    p.append(p_pr)

    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.extend([fonts, size])
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    tab_run = OxmlElement("w:r")
    tab_run.append(OxmlElement("w:tab"))
    p.extend([run, tab_run, field_run(f"PAGEREF {bookmark} \\h", "?")])
    return p


def find_placeholder_after(doc, heading_text):
    paragraphs = doc.paragraphs
    for idx, paragraph in enumerate(paragraphs):
        if paragraph.text.strip() == heading_text:
            for candidate in paragraphs[idx + 1 :]:
                if candidate.style and candidate.style.name == "Heading 1":
                    break
                if candidate.text.strip() == "No table of contents entries found.":
                    return candidate
    raise RuntimeError(f"Placeholder untuk {heading_text} tidak ditemukan")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("docx")
    args = parser.parse_args()
    path = Path(args.docx).resolve()
    doc = Document(path)

    figures = [p for p in doc.paragraphs if p.style and p.style.name == "Figure Caption"]
    tables = [p for p in doc.paragraphs if p.style and p.style.name == "Table Caption"]
    for idx, paragraph in enumerate(figures, 1):
        add_bookmark(paragraph, f"figcap{idx}", 1000 + idx)
    for idx, paragraph in enumerate(tables, 1):
        add_bookmark(paragraph, f"tblcap{idx}", 2000 + idx)

    for heading_text, entries, prefix in (
        ("DAFTAR GAMBAR", figures, "figcap"),
        ("DAFTAR TABEL", tables, "tblcap"),
    ):
        placeholder = find_placeholder_after(doc, heading_text)
        anchor = placeholder._p
        for idx, caption in enumerate(entries, 1):
            new_p = index_paragraph(caption.text.strip(), f"{prefix}{idx}")
            anchor.addnext(new_p)
            anchor = new_p
        placeholder._element.getparent().remove(placeholder._element)

    doc.save(path)
    print(f"PATCHED={path}")
    print(f"FIGURES={len(figures)} TABLES={len(tables)}")


if __name__ == "__main__":
    main()
